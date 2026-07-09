from fastapi.testclient import TestClient
from io import BytesIO
from docx import Document
from sqlmodel import SQLModel

from app.database import engine, normalize_generated_recruiter_source
from app.main import app


def reset_database() -> None:
    SQLModel.metadata.drop_all(engine)
    SQLModel.metadata.create_all(engine)


def test_health_returns_ok():
    client = TestClient(app)
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok", "service": "job-search-agent"}


def test_root_returns_api_entrypoint():
    client = TestClient(app)
    response = client.get("/")
    assert response.status_code == 200
    assert response.json() == {
        "service": "job-search-agent",
        "health": "/api/health",
        "docs": "/docs",
    }


def test_create_vacancy_and_list_it():
    reset_database()
    with TestClient(app) as client:
        response = client.post(
            "/api/vacancies",
            json={"external_id": "MAN-1", "company": "TestCo", "title": "Lead Product Designer"},
        )
        assert response.status_code == 200
        created = response.json()
        assert created["company"] == "TestCo"

        list_response = client.get("/api/vacancies")
        assert list_response.status_code == 200
        assert any(row["external_id"] == "MAN-1" for row in list_response.json())


def test_create_vacancy_from_url_extracts_source_fields(monkeypatch):
    reset_database()

    class Extracted:
        source = "HH.ru"
        company = "Product Studio"
        title = "Senior Product Designer"
        location = "Remote"
        language = "English"
        description_raw = "Lead UX strategy, Figma components and design systems for product teams."

    monkeypatch.setattr("app.api.routes.extract_vacancy_from_url", lambda _url: Extracted())
    with TestClient(app) as client:
        response = client.post(
            "/api/vacancies",
            json={"source_url": "https://hh.ru/vacancy/555"},
        )

        assert response.status_code == 200
        created = response.json()
        assert created["source"] == "HH.ru"
        assert created["company"] == "Product Studio"
        assert created["title"] == "Senior Product Designer"
        assert created["location"] == "Remote"
        assert created["language"] == "English"
        assert created["description_raw"].startswith("Lead UX strategy")
        assert "Design Systems" in created["vacancy_keywords"]


def test_create_vacancy_accepts_source_description_and_generates_tailored_cv():
    reset_database()
    with TestClient(app) as client:
        candidate = client.post(
            "/api/candidate/text",
            json={"text": "Aleksandr Grenkov Lead Product Designer Automotive UX Voice UX Design Systems English B2"},
        )
        assert candidate.status_code == 200

        created_response = client.post(
            "/api/vacancies",
            json={
                "external_id": "REAL-1",
                "source": "LinkedIn",
                "company": "Rivian",
                "title": "Sr. Lead Product Designer - Design System Frameworks",
                "location": "Palo Alto / Remote",
                "language": "English",
                "source_url": "https://www.linkedin.com/jobs/example",
                "description_raw": "Design system frameworks for automotive HMI products. Partner with engineering and product.",
                "requirements": "Design systems; Automotive UX; HMI; Figma; component governance",
                "responsibilities": "Define reusable frameworks and improve product consistency across vehicle experiences.",
            },
        )
        assert created_response.status_code == 200
        created = created_response.json()
        assert created["description_raw"].startswith("Design system frameworks")
        assert "Design Systems" in created["vacancy_keywords"]
        assert "Automotive UX" in created["vacancy_keywords"]
        assert "Component Governance" in created["vacancy_keywords"]

        analysis_response = client.post("/api/analysis/run")
        assert analysis_response.status_code == 200

        vacancy = client.get("/api/vacancies").json()[0]
        assert vacancy["cv_file_path"] == f"/api/vacancies/{vacancy['id']}/tailored-cv.docx"
        assert "Design Systems" in vacancy["top_match_keywords"]
        assert "Automotive" in vacancy["top_match_keywords"]

        cv_response = client.get(vacancy["cv_file_path"])
        assert cv_response.status_code == 200
        document = Document(BytesIO(cv_response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "Rivian" in text
        assert "Sr. Lead Product Designer - Design System Frameworks" in text
        assert "Design Systems" in text
        assert "Automotive UX" in text
        assert "EXECUTIVE SUMMARY" in text
        assert "SELECTED CAREER IMPACT" in text
        assert "CORE EXPERTISE" in text
        assert "EXPERIENCE" in text
        assert "Key achievements" in text
        assert "EDUCATION" in text
        assert "TOOLS" in text
        assert "CERTIFICATIONS" in text
        assert "LANGUAGES" in text
        assert "Led end-to-end product design" in text
        assert "Tailored for:" not in text
        assert "Fit:" not in text
        assert "Vacancy Requirements Mirrored" not in text
        assert "general" not in text
        assert "No unverified metrics" not in text
        assert "Adapt this" not in text
        assert "Worked on" not in text
        assert "Responsible for" not in text
        assert "Participated in" not in text
        assert "Helped" not in text
        assert "Vacancy Source" not in text
        assert "Vacancy Link" not in text
        assert "https://www.linkedin.com/jobs/example" not in text


def test_tailored_cv_header_contacts_expertise_tags_and_no_work_dates():
    reset_database()
    cv_text = """Aleksandr Grenkov
Lead Product Designer
aleksandr@example.com
https://www.linkedin.com/in/aleksandr-grenkov
https://grenkov.design
Telegram: @agrenkov
Automotive UX Voice UX Design Systems Enterprise UX English B2
"""
    with TestClient(app) as client:
        candidate = client.post("/api/candidate/text", json={"text": cv_text})
        assert candidate.status_code == 200

        created_response = client.post(
            "/api/vacancies",
            json={
                "external_id": "CONTACTS-1",
                "source": "LinkedIn",
                "company": "VehicleCo",
                "title": "Lead Product Designer - Automotive HMI",
                "location": "Remote",
                "language": "English",
                "description_raw": "Lead UX for automotive HMI, design systems, and multimodal interaction.",
                "requirements": "Automotive UX; HMI; Voice UX; Figma; Design Systems; Component Governance",
                "responsibilities": "Define product strategy and partner with engineering on vehicle interfaces.",
            },
        )
        assert created_response.status_code == 200
        assert client.post("/api/analysis/run").status_code == 200

        vacancy = client.get("/api/vacancies").json()[0]
        cv_response = client.get(vacancy["cv_file_path"])
        assert cv_response.status_code == 200

        document = Document(BytesIO(cv_response.content))
        paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
        text = "\n".join(paragraph.text for paragraph in paragraphs)

        assert "ALEKSANDR GRENKOV" in text
        assert "Candidate" not in text
        assert "aleksandr@example.com" in text
        assert "linkedin.com/in/aleksandr-grenkov" in text
        assert "grenkov.design" in text
        assert "@agrenkov" in text
        assert "Open to international and remote opportunities." in text

        core_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text == "CORE EXPERTISE")
        core_tags = paragraphs[core_index + 1]
        assert core_tags.style.name != "List Bullet"
        assert " • " in core_tags.text
        assert "Automotive UX" in core_tags.text
        assert "HMI" in core_tags.text
        assert "Voice UX" in core_tags.text
        assert "Design Systems" in core_tags.text

        assert "2023-2026" not in text
        assert "2022-2023" not in text


def test_candidate_contacts_endpoint_adds_missing_contacts_to_tailored_cv():
    reset_database()
    with TestClient(app) as client:
        candidate = client.post(
            "/api/candidate/text",
            json={"text": "Aleksandr Grenkov Lead Product Designer Automotive UX Voice UX Design Systems English B2"},
        )
        assert candidate.status_code == 200

        contacts = client.post(
            "/api/candidate/contacts",
            json={
                "email": "aleksandr@example.com",
                "linkedin": "https://www.linkedin.com/in/aleksandr-grenkov",
                "portfolio": "https://grenkov.design",
                "telegram": "@agrenkov",
            },
        )
        assert contacts.status_code == 200
        assert "aleksandr@example.com" in contacts.json()["raw_cv_text"]

        created_response = client.post(
            "/api/vacancies",
            json={
                "external_id": "PATCH-CONTACTS-1",
                "source": "LinkedIn",
                "company": "VehicleCo",
                "title": "Lead Product Designer",
                "language": "English",
                "requirements": "Automotive UX; Voice UX; Design Systems",
            },
        )
        assert created_response.status_code == 200
        assert client.post("/api/analysis/run").status_code == 200

        vacancy = client.get("/api/vacancies").json()[0]
        cv_response = client.get(vacancy["cv_file_path"])
        assert cv_response.status_code == 200
        document = Document(BytesIO(cv_response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        assert "aleksandr@example.com" in text
        assert "linkedin.com/in/aleksandr-grenkov" in text
        assert "grenkov.design" in text
        assert "@agrenkov" in text


def test_tailored_cv_uses_russian_for_russian_vacancy():
    reset_database()
    with TestClient(app) as client:
        client.post(
            "/api/candidate/text",
            json={"text": "Aleksandr Grenkov Lead Product Designer Automotive UX HMI Voice UX Design Systems English B2"},
        )
        created_response = client.post(
            "/api/vacancies",
            json={
                "external_id": "RU-1",
                "source": "Telegram @wantapply_design",
                "company": "Команда продукта",
                "title": "Ведущий продуктовый дизайнер",
                "language": "Russian",
                "description_raw": "Ищем продуктового дизайнера для дизайн-системы, интерфейсов и сложных пользовательских сценариев.",
                "requirements": "Figma; дизайн-системы; UX; интерфейсы",
                "responsibilities": "Проектировать продуктовые сценарии и развивать дизайн-систему.",
            },
        )
        assert created_response.status_code == 200
        assert client.post("/api/analysis/run").status_code == 200

        vacancy = client.get("/api/vacancies").json()[0]
        cv_response = client.get(vacancy["cv_file_path"])
        assert cv_response.status_code == 200

        document = Document(BytesIO(cv_response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "ПРОФЕССИОНАЛЬНЫЙ ПРОФИЛЬ" in text
        assert "КЛЮЧЕВОЙ КАРЬЕРНЫЙ ЭФФЕКТ" in text
        assert "КЛЮЧЕВАЯ ЭКСПЕРТИЗА" in text
        assert "ОПЫТ" in text
        assert "Ключевые достижения" in text
        assert "ОБРАЗОВАНИЕ" in text
        assert "ИНСТРУМЕНТЫ" in text
        assert "СЕРТИФИКАЦИИ" in text
        assert "ЯЗЫКИ" in text
        assert "Адаптировано под:" not in text
        assert "Соответствие:" not in text
        assert "общее" not in text
        assert "Без неподтвержденных метрик" not in text
        assert "Profile" not in text
        assert "Skills" not in text
        assert "Experience" not in text


def test_tailored_cv_uses_vacancy_text_language_over_metadata():
    reset_database()
    with TestClient(app) as client:
        client.post(
            "/api/candidate/text",
            json={"text": "Aleksandr Grenkov Lead Product Designer Automotive UX HMI Voice UX Design Systems English B2"},
        )
        created_response = client.post(
            "/api/vacancies",
            json={
                "external_id": "EN-META-RU",
                "source": "LinkedIn",
                "company": "VehicleCo",
                "title": "Senior Product Designer",
                "language": "Russian",
                "description_raw": "We are looking for a Senior Product Designer to build design systems, dashboards, and complex UX flows.",
                "requirements": "Figma; design systems; UX research; English",
                "responsibilities": "Lead product design discovery and partner with product and engineering.",
            },
        )
        assert created_response.status_code == 200
        assert client.post("/api/analysis/run").status_code == 200

        vacancy = client.get("/api/vacancies").json()[0]
        cv_response = client.get(vacancy["cv_file_path"])
        assert cv_response.status_code == 200

        document = Document(BytesIO(cv_response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "EXECUTIVE SUMMARY" in text
        assert "CORE EXPERTISE" in text
        assert "EXPERIENCE" in text
        assert "Профиль" not in text
        assert "Навыки" not in text
        assert "Опыт" not in text


def test_status_update_rejects_caller_controlled_actor():
    reset_database()
    with TestClient(app) as client:
        created = client.post(
            "/api/vacancies",
            json={"external_id": "MAN-2", "company": "TestCo", "title": "Lead Product Designer"},
        ).json()
        response = client.patch(
            f"/api/vacancies/{created['id']}/status",
            json={"status": "Sent", "actor": "system"},
        )
        assert response.status_code == 422


def test_user_status_update_can_mark_sent():
    reset_database()
    with TestClient(app) as client:
        created = client.post(
            "/api/vacancies",
            json={"external_id": "MAN-3", "company": "TestCo", "title": "Lead Product Designer"},
        ).json()
        response = client.patch(
            f"/api/vacancies/{created['id']}/status",
            json={"status": "Sent"},
        )
        assert response.status_code == 200
        assert response.json()["submit_status"] == "Sent"


def test_candidate_analysis_materials_and_export_flow():
    reset_database()
    with TestClient(app) as client:
        candidate_response = client.post(
            "/api/candidate/text",
            json={"text": "Aleksandr Grenkov Lead Product Designer Automotive UX Voice UX English B2"},
        )
        assert candidate_response.status_code == 200

        vacancy_response = client.post(
            "/api/vacancies",
            json={
                "external_id": "L08",
                "company": "42dot",
                "title": "Lead Brand / UI Designer - Automotive",
                "language": "English",
            },
        )
        assert vacancy_response.status_code == 200

        analysis_response = client.post("/api/analysis/run")
        assert analysis_response.status_code == 200
        assert analysis_response.json() == {"analyzed": 1}

        materials_response = client.get("/api/materials")
        assert materials_response.status_code == 200
        materials = materials_response.json()
        assert len(materials) == 1
        assert "42dot" in materials[0]["short_note"]

        export_response = client.get("/api/exports/zip")
        assert export_response.status_code == 200
        assert export_response.headers["content-type"] == "application/zip"


def test_analysis_from_cv_requires_candidate_profile():
    reset_database()
    with TestClient(app) as client:
        response = client.post("/api/analysis/from-cv")
        assert response.status_code == 400
        assert response.json()["detail"] == "Create a candidate profile first"


def test_analysis_from_cv_generates_twenty_recruiter_role_matches():
    reset_database()
    cv_text = (
        "Aleksandr Grenkov Lead Product Designer Automotive UX HMI Voice UX "
        "Conversational Design Design Systems Smart City Transport Enterprise UX "
        "Telecom B2B B2C Mobile UX English B2"
    )
    with TestClient(app) as client:
        candidate_response = client.post("/api/candidate/text", json={"text": cv_text})
        assert candidate_response.status_code == 200

        response = client.post("/api/analysis/from-cv")
        assert response.status_code == 200
        assert response.json() == {"generated": 20, "analyzed": 20}

        vacancies = client.get("/api/vacancies").json()
        assert len(vacancies) == 20
        assert vacancies[0]["source"] == "LinkedIn"
        assert all(row["source"] != "CV Recruiter Match" for row in vacancies)
        assert vacancies[0]["company"] == "Target role"
        assert vacancies[0]["date_status"] == "generated from CV, not a live vacancy"
        assert vacancies[0]["title"] == "Lead Product Designer"
        assert "Product Design" in vacancies[0]["top_match_keywords"]
        assert "English B2" in vacancies[0]["gaps_risks"]
        assert {row["rank"] for row in vacancies} == set(range(1, 21))

        materials = client.get("/api/materials").json()
        assert len(materials) == 20
        assert "Lead Product Designer" in materials[0]["short_note"]


def test_analysis_from_sources_collects_real_source_vacancies(monkeypatch):
    reset_database()
    cv_text = "Aleksandr Grenkov Lead Product Designer Automotive UX HMI Design Systems English B2"

    def fake_collect_live_vacancies(profile, roles):
        assert "Lead Product Designer" in profile.raw_cv_text
        assert roles
        from app.models import Vacancy

        return [
            Vacancy(
                external_id="HH-123",
                source="HH.ru",
                company="AutoTech",
                title="Lead Product Designer HMI",
                source_url="https://hh.ru/vacancy/123",
                description_raw="Design systems and HMI",
                date_status="verified within 7 days",
            ),
            Vacancy(
                external_id="TG-wantapply_design-77",
                source="Telegram @wantapply_design",
                company="@wantapply_design",
                title="Senior Product Designer",
                source_url="https://t.me/wantapply_design/77",
                description_raw="Remote product design vacancy",
                date_status="verified within 7 days",
            ),
        ]

    monkeypatch.setattr("app.api.routes.collect_live_vacancies", fake_collect_live_vacancies)
    with TestClient(app) as client:
        candidate_response = client.post("/api/candidate/text", json={"text": cv_text})
        assert candidate_response.status_code == 200

        response = client.post("/api/analysis/from-sources")
        assert response.status_code == 200
        assert response.json() == {"generated": 2, "analyzed": 2, "fallback": False}

        vacancies = client.get("/api/vacancies").json()
        assert [row["source"] for row in vacancies] == ["HH.ru", "Telegram @wantapply_design"]
        assert all(row["cv_file_path"] == f"/api/vacancies/{row['id']}/tailored-cv.docx" for row in vacancies)


def test_analysis_from_sources_does_not_create_search_link_fallback(monkeypatch):
    reset_database()
    monkeypatch.setattr("app.api.routes.collect_live_vacancies", lambda _profile, _roles: [])
    with TestClient(app) as client:
        client.post("/api/candidate/text", json={"text": "Aleksandr Grenkov Lead Product Designer"})
        client.post(
            "/api/vacancies",
            json={
                "external_id": "OLD-SEARCH",
                "source": "LinkedIn",
                "company": "Target role",
                "title": "Lead Product Designer",
                "source_url": "https://www.linkedin.com/jobs/search/?keywords=Lead+Product+Designer",
            },
        )

        response = client.post("/api/analysis/from-sources")
        assert response.status_code == 200
        assert response.json() == {"generated": 0, "analyzed": 0, "fallback": False}
        assert client.get("/api/vacancies").json() == []


def test_startup_normalizes_legacy_recruiter_match_source():
    reset_database()
    with TestClient(app) as client:
        response = client.post(
            "/api/vacancies",
            json={
                "external_id": "LEGACY-1",
                "source": "CV Recruiter Match",
                "company": "Target role",
                "title": "Lead Product Designer",
                "source_url": "https://www.linkedin.com/jobs/search/?keywords=Lead+Product+Designer",
            },
        )
        assert response.status_code == 200

        normalize_generated_recruiter_source(engine)

        vacancies = client.get("/api/vacancies").json()
        assert vacancies[0]["source"] == "LinkedIn"


def test_analysis_from_cv_links_every_row_to_tailored_cv_docx():
    reset_database()
    cv_text = (
        "Aleksandr Grenkov Lead Product Designer Automotive UX HMI Voice UX "
        "Design Systems Smart City Transport Enterprise UX English B2"
    )
    with TestClient(app) as client:
        client.post("/api/candidate/text", json={"text": cv_text})
        client.post("/api/analysis/from-cv")

        vacancies = client.get("/api/vacancies").json()
        assert len(vacancies) == 20
        assert all(row["cv_file_path"] == f"/api/vacancies/{row['id']}/tailored-cv.docx" for row in vacancies)

        first = vacancies[0]
        response = client.get(first["cv_file_path"])
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")

        document = Document(BytesIO(response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "ALEKSANDR GRENKOV" in text
        assert first["title"] in text
        assert "EXECUTIVE SUMMARY" in text
        assert "SELECTED CAREER IMPACT" in text
        assert "CORE EXPERTISE" in text
        assert "EXPERIENCE" in text
        assert "Key achievements" in text
        assert "EDUCATION" in text
        assert "TOOLS" in text
        assert "CERTIFICATIONS" in text
        assert "LANGUAGES" in text
        assert "Exact Match Keywords" not in text
        assert "Google XYZ Tailored Bullets" not in text
        assert "Selected Experience" not in text
        assert "Adaptation Strategy" not in text
        assert "Recruiter-Safe Notes" not in text
        assert "Tailored for:" not in text
        assert "Fit:" not in text
        assert "No unverified metrics" not in text
        assert "Adapt this" not in text
        assert "Worked on" not in text
        assert "Responsible for" not in text
        assert "Participated in" not in text
        assert "Helped" not in text
        assert "Vacancy Source" not in text
        assert "Vacancy Link" not in text
        assert "English B2" in text
        assert "English C1" not in text
        assert "increased conversion by" not in text.lower()


def test_master_cv_template_docx_is_recruiter_safe_and_adaptable():
    reset_database()
    with TestClient(app) as client:
        client.post(
            "/api/candidate/text",
            json={
                "text": (
                    "Aleksandr Grenkov Lead Product Designer Automotive UX HMI Voice UX "
                    "Design Systems Smart City Transport Enterprise UX English B2"
                )
            },
        )
        response = client.get("/api/candidate/master-cv.docx")
        assert response.status_code == 200

        document = Document(BytesIO(response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "ATS-OPTIMIZED MASTER CV TEMPLATE" in text
        assert len(document.tables) == 0
        for section in [
            "Contact",
            "Target Title",
            "Professional Summary",
            "Target Keywords",
            "Core Skills",
            "Domain Expertise",
            "Professional Experience",
            "Selected Projects",
            "Education",
            "Languages",
            "Tools",
        ]:
            assert section in text
        assert "[TARGET ROLE]" in text
        assert "[ROLE]" in text
        assert "[COMPANY]" in text
        assert "[DOMAIN KEYWORDS]" in text
        assert "[VACANCY KEYWORDS]" in text
        assert "Google XYZ" in text
        assert "English B2" in text
        assert "English C1" not in text
        assert "worked on" not in text.lower()


def test_candidate_upload_sanitizes_filename(tmp_path, monkeypatch):
    reset_database()
    from app.config import get_settings

    settings = get_settings()
    monkeypatch.setattr(settings, "storage_dir", tmp_path)
    with TestClient(app) as client:
        response = client.post(
            "/api/candidate/upload",
            files={"file": ("../cv.txt", b"Aleksandr Grenkov Lead Product Designer", "text/plain")},
        )
        assert response.status_code == 200
    assert (tmp_path / "uploads" / "cv.txt").exists()
    assert not (tmp_path / "cv.txt").exists()


def test_upload_vacancy_file_imports_csv_rows():
    reset_database()
    csv_bytes = (
        "Rank,ID,Source,Company,Vacancy,Location,Posted,Date Status,Language,Fit,Priority,Submit Status,Next Action,CV File Path,Source URL,Tailored Headline,Top Match Keywords,Gaps / Risks,Adaptation Strategy\n"
        "1,L08,LinkedIn,42dot,Lead Brand / UI Designer - Automotive,Remote,2 days ago,verified,English,96,Very High,Ready to submit manually,Submit manually,Tailored_CVs/08.docx,https://example.com/42dot,Lead Product Designer,Automotive; HMI,US location risk,Lead with ATOM\n"
        "2,L04,LinkedIn,ZOE,Lead Product Designer - Design System,Remote,4 days ago,verified,English,95,Very High,Ready to submit manually,Submit manually,Tailored_CVs/04.docx,https://example.com/zoe,Design Systems Lead,Design systems,Need token examples,Lead with systems\n"
    ).encode("utf-8")
    with TestClient(app) as client:
        response = client.post(
            "/api/imports/vacancies/upload",
            files={"file": ("vacancies.csv", csv_bytes, "text/csv")},
        )
        assert response.status_code == 200
        assert response.json() == {"imported": 2}

        vacancies_response = client.get("/api/vacancies")
        assert vacancies_response.status_code == 200
        rows = vacancies_response.json()
        assert [row["company"] for row in rows] == ["42dot", "ZOE"]
        assert rows[0]["source_url"] == "https://example.com/42dot"


def test_upload_vacancy_file_rejects_unsupported_format():
    reset_database()
    with TestClient(app) as client:
        response = client.post(
            "/api/imports/vacancies/upload",
            files={"file": ("vacancies.json", b"{}", "application/json")},
        )
        assert response.status_code == 400
        assert "Unsupported vacancy file type" in response.json()["detail"]
