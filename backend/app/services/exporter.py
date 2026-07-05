import csv
import html
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook

from app.models import ApplicationMaterial, Vacancy


def export_queue_csv(output_dir: Path, vacancies: list[Vacancy]) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "application_queue.csv"
    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow([
            "ID",
            "Source",
            "Company",
            "Vacancy",
            "Priority",
            "Fit",
            "Status",
            "CV File Path",
            "Source URL",
            "Vacancy Keywords",
            "Vacancy Source Text",
            "Requirements",
            "Responsibilities",
        ])
        for vacancy in vacancies:
            writer.writerow(
                [
                    vacancy.external_id,
                    vacancy.source,
                    vacancy.company,
                    vacancy.title,
                    vacancy.priority,
                    vacancy.fit_score,
                    vacancy.submit_status,
                    vacancy.cv_file_path,
                    vacancy.source_url,
                    vacancy.vacancy_keywords,
                    vacancy.description_raw,
                    vacancy.requirements,
                    vacancy.responsibilities,
                ]
            )
    return path


def export_queue_html(output_dir: Path, vacancies: list[Vacancy]) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "cv_links_index.html"
    rows = []
    for vacancy in vacancies:
        cv_link = f'<a href="{html.escape(vacancy.cv_file_path)}">Open CV</a>' if vacancy.cv_file_path else ""
        source_link = f'<a href="{html.escape(vacancy.source_url)}">Open vacancy</a>' if vacancy.source_url else ""
        rows.append(
            "<tr>"
            f"<td>{html.escape(vacancy.external_id)}</td>"
            f"<td>{html.escape(vacancy.source)}</td>"
            f"<td>{html.escape(vacancy.company)}</td>"
            f"<td>{html.escape(vacancy.title)}</td>"
            f"<td>{html.escape(vacancy.priority)}</td>"
            f"<td>{cv_link}</td>"
            f"<td>{source_link}</td>"
            "</tr>"
        )
    path.write_text(
        '<!doctype html><html><head><meta charset="utf-8"><title>Ready-to-send applications</title></head>'
        "<body><h1>Ready-to-send applications</h1><table>"
        "<thead><tr><th>ID</th><th>Source</th><th>Company</th><th>Vacancy</th><th>Priority</th><th>CV</th><th>Vacancy</th></tr></thead>"
        f"<tbody>{''.join(rows)}</tbody></table></body></html>",
        encoding="utf-8",
    )
    return path


def export_queue_xlsx(output_dir: Path, vacancies: list[Vacancy]) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "application_queue.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Application Queue"
    sheet.append([
        "ID",
        "Source",
        "Company",
        "Vacancy",
        "Priority",
        "Fit",
        "Status",
        "CV File Path",
        "Source URL",
        "Vacancy Keywords",
        "Vacancy Source Text",
        "Requirements",
        "Responsibilities",
    ])
    for vacancy in vacancies:
        sheet.append(
            [
                vacancy.external_id,
                vacancy.source,
                vacancy.company,
                vacancy.title,
                vacancy.priority,
                vacancy.fit_score,
                str(vacancy.submit_status),
                vacancy.cv_file_path,
                vacancy.source_url,
                vacancy.vacancy_keywords,
                vacancy.description_raw,
                vacancy.requirements,
                vacancy.responsibilities,
            ]
        )
        row_number = sheet.max_row
        if vacancy.cv_file_path:
            sheet.cell(row=row_number, column=8).hyperlink = vacancy.cv_file_path
            sheet.cell(row=row_number, column=8).style = "Hyperlink"
        if vacancy.source_url:
            sheet.cell(row=row_number, column=9).hyperlink = vacancy.source_url
            sheet.cell(row=row_number, column=9).style = "Hyperlink"
    for column in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 12), 60)
    workbook.save(path)
    return path


def export_zip_package(output_dir: Path, vacancies: list[Vacancy], materials: list[ApplicationMaterial]) -> Path:
    csv_path = export_queue_csv(output_dir, vacancies)
    xlsx_path = export_queue_xlsx(output_dir, vacancies)
    html_path = export_queue_html(output_dir, vacancies)
    notes_path = output_dir / "ready_to_send_messages.txt"
    material_by_vacancy = {material.vacancy_id: material for material in materials}
    notes = []
    for vacancy in vacancies:
        material = material_by_vacancy.get(vacancy.id or 0)
        if material:
            notes.append(f"{vacancy.external_id} - {vacancy.company}\n{material.email_cover_letter}\n")
    notes_path.write_text("\n---\n".join(notes), encoding="utf-8")
    zip_path = output_dir / "ready_to_send_package.zip"
    with ZipFile(zip_path, "w") as archive:
        archive.write(csv_path, "application_queue.csv")
        archive.write(xlsx_path, "application_queue.xlsx")
        archive.write(html_path, "cv_links_index.html")
        archive.write(notes_path, "ready_to_send_messages.txt")
    return zip_path
