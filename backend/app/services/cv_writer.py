from io import BytesIO
import re
from typing import Optional

from docx import Document
from docx.shared import Inches, Pt

from app.models import CandidateProfile, Vacancy


SAFE_ACTION_VERBS = ("Led", "Defined", "Established", "Structured", "Simplified", "Partnered", "Standardized")


def build_master_cv_document(profile: CandidateProfile) -> BytesIO:
    document = _base_document()
    _add_title(document, "ALEKSANDR GRENKOV", "ATS-OPTIMIZED MASTER CV TEMPLATE")
    document.add_heading("Contact", level=2)
    document.add_paragraph("Email | LinkedIn | Portfolio | Telegram | Location/relocation: [TARGET LOCATION]")
    document.add_heading("Target Title", level=2)
    document.add_paragraph("[TARGET ROLE] / [ROLE] at [COMPANY]")
    document.add_heading("Professional Summary", level=2)
    document.add_paragraph(
        "Lead Product Designer with 10+ years across product design, UX strategy, design systems, "
        "telecom, e-commerce, automotive UX, smart city systems, enterprise interfaces, voice UX, "
        "sound design, and AI-related interaction design. Adapt this summary with [DOMAIN KEYWORDS] "
        "and [VACANCY KEYWORDS] while keeping all claims evidence-based."
    )
    document.add_heading("Target Keywords", level=2)
    document.add_paragraph(
        "[VACANCY KEYWORDS]; [DOMAIN KEYWORDS]; Product Design; UX Strategy; UX/UI Design; "
        "Design Systems; Voice UX; Conversational Design; Automotive UX; HMI; Enterprise UX; "
        "B2B; B2C; Mobile UX; Design Leadership"
    )
    document.add_heading("Core Skills", level=2)
    document.add_paragraph(
        "Product Design; UX Strategy; Information Architecture; User Flows; Prototyping; "
        "Design Systems; Design Review; UX Research; CJM; Mobile UX; Dashboard UX; "
        "Cross-functional Collaboration; Mentoring"
    )
    document.add_heading("Domain Expertise", level=2)
    document.add_paragraph(
        "Automotive UX; HMI; In-vehicle UX; Voice Assistant UX; Sound Identity; AVAS; "
        "Smart City UX; Transport Systems; Parking Management; Telecom Products; "
        "Enterprise Platforms; B2B/B2C Digital Products"
    )
    document.add_heading("Google XYZ Bullet Formula", level=2)
    document.add_paragraph("Use: Accomplished [X], as measured by [Y], by doing [Z].")
    document.add_paragraph("If [Y] is not verified, do not invent metrics. Use truthful impact-oriented wording.")
    document.add_heading("Professional Experience", level=2)
    _add_experience_template(
        document,
        "Lead Product Designer — ATOM, Electric Vehicle Startup",
        [
            "Established a scalable voice interaction framework by defining intents, prompts, tone of voice, and multimodal patterns for in-car assistant scenarios.",
            "Structured in-vehicle UX logic by mapping customer journeys, information architecture, and reusable interaction patterns for HMI scenarios.",
            "Defined product UX for sound identity, AVAS, and internal sound indication systems by partnering with engineering and research teams.",
        ],
    )
    _add_experience_template(
        document,
        "Product Designer — Information Technology Factory",
        [
            "Simplified city-scale operator workflows by designing dashboards, parking apps, and transport-system interfaces for monitoring and control scenarios.",
            "Delivered mobile parking app experiences by translating research, user flows, and information architecture into Android and iOS interfaces.",
        ],
    )
    _add_experience_template(
        document,
        "Lead Product Designer — VEON / Beeline Russia",
        [
            "Standardized complex product interfaces by developing design-system practices, reusable UI patterns, and cross-functional design review rituals.",
            "Partnered with product, engineering, research, and external design teams to translate ambiguous requirements into clear user flows and production-ready interfaces.",
        ],
    )
    _add_experience_template(
        document,
        "Product Designer — ABBYY Software House",
        [
            "Delivered UI/UX for international software products by structuring product flows and interface patterns for Lingvo and PDF Transformer scenarios.",
        ],
    )
    document.add_heading("Selected Projects", level=2)
    _add_bullets(
        document,
        [
            "In-car voice assistant and HMI interaction logic — adapt with [VACANCY KEYWORDS] for automotive, AI, voice, or HMI roles.",
            "Smart city and transport dashboards — adapt with [DOMAIN KEYWORDS] for enterprise, mobility, transport, or operator UX roles.",
            "Telecom B2B/B2C product and design-system work — adapt for design systems, platform UX, and complex product roles.",
        ],
    )
    document.add_heading("Education", level=2)
    document.add_paragraph("[EDUCATION] — keep factual, short, and ATS-readable.")
    document.add_heading("Languages", level=2)
    document.add_paragraph("Russian: native; English B2")
    document.add_heading("Tools", level=2)
    document.add_paragraph("Figma; Prototyping; Design Systems; UX Research; Information Architecture; CJM; Jira; Confluence; Adobe tools")
    document.add_heading("Red Flag Guardrails", level=2)
    _add_bullets(
        document,
        [
            "English level: English B2. Keep language claims at the verified level.",
            "Do not invent conversion, revenue, launch, budget, or team-size metrics.",
            "Do not claim ML engineering ownership; describe AI work as product, UX, prompts, intents, and interaction design.",
            "Replace weak phrasing with specific action verbs: " + ", ".join(SAFE_ACTION_VERBS) + ".",
        ],
    )
    document.add_heading("Source CV Constraints", level=2)
    document.add_paragraph(profile.constraints)
    return _save(document)


def build_tailored_cv_document(profile: CandidateProfile, vacancy: Vacancy) -> BytesIO:
    if _is_russian_vacancy(vacancy):
        return _build_russian_tailored_resume_document(profile, vacancy)
    return _build_english_tailored_resume_document(profile, vacancy)


def _build_english_tailored_resume_document(profile: CandidateProfile, vacancy: Vacancy) -> BytesIO:
    document = _base_document()
    _add_title(document, _candidate_display_name(profile), vacancy.title or "Lead Product Designer", profile=profile)
    document.add_paragraph(_headline_tags(vacancy))

    document.add_heading("EXECUTIVE SUMMARY", level=2)
    document.add_paragraph(
        f"Lead Product Designer with 10+ years of experience delivering digital products aligned with {vacancy.title}, "
        f"with deep expertise in {_domain_focus(vacancy)}."
    )
    document.add_paragraph(
        "Expert in transforming complex technical systems into intuitive customer experiences through Product Strategy, "
        "UX Leadership, Design Systems, Human-Centered AI and cross-functional product delivery."
    )
    document.add_paragraph(
        f"Led end-to-end product design for voice assistants, automotive HMI, enterprise platforms and consumer products "
        f"while partnering with Product Managers, Engineering, AI, Research and Executive Leadership for {vacancy.company or 'product teams'}."
    )

    document.add_heading("CORE EXPERTISE", level=2)
    _add_tag_paragraph(document, _core_expertise(vacancy))

    document.add_heading("SELECTED CAREER IMPACT", level=2)
    _add_bullets(document, _career_impact_bullets(vacancy))

    document.add_heading("EXPERIENCE", level=2)
    for heading, description, bullets in _resume_experience_blocks(vacancy):
        paragraph = document.add_paragraph()
        paragraph.add_run(heading).bold = True
        document.add_paragraph(description)
        key = document.add_paragraph()
        key.add_run("Key achievements").bold = True
        _add_bullets(document, bullets)

    document.add_heading("EDUCATION", level=2)
    document.add_paragraph("Education details available upon request.")
    document.add_heading("TOOLS", level=2)
    document.add_paragraph("Figma; Miro; Jira; Confluence; Adobe Creative Suite; Prototyping; Design Systems; UX Research")
    document.add_heading("CERTIFICATIONS", level=2)
    document.add_paragraph("Relevant product design, UX and leadership certifications available upon request.")
    document.add_heading("LANGUAGES", level=2)
    document.add_paragraph(profile.languages or "Russian native; English B2")
    return _save(document)


def _build_russian_tailored_resume_document(profile: CandidateProfile, vacancy: Vacancy) -> BytesIO:
    document = _base_document()
    _add_title(
        document,
        _candidate_display_name(profile),
        vacancy.title or "Ведущий продуктовый дизайнер",
        language="ru",
        profile=profile,
    )
    document.add_paragraph(_headline_tags(vacancy, language="ru"))

    document.add_heading("ПРОФЕССИОНАЛЬНЫЙ ПРОФИЛЬ", level=2)
    document.add_paragraph(
        f"Ведущий продуктовый дизайнер с опытом 10+ лет в создании цифровых продуктов под задачи роли «{vacancy.title}», "
        f"с сильной экспертизой в направлениях: {_domain_focus(vacancy, language='ru')}."
    )
    document.add_paragraph(
        "Сильная сторона — превращать сложные технические системы в понятный клиентский опыт через продуктовую стратегию, "
        "UX-лидерство, дизайн-системы, Human-Centered AI и системное взаимодействие с командами продукта и разработки."
    )
    document.add_paragraph(
        f"Вел end-to-end продуктовый дизайн голосовых ассистентов, автомобильных HMI, enterprise-платформ и consumer-продуктов "
        f"во взаимодействии с Product Management, Engineering, AI, Research и руководством."
    )

    document.add_heading("КЛЮЧЕВАЯ ЭКСПЕРТИЗА", level=2)
    _add_tag_paragraph(document, _core_expertise(vacancy, language="ru"))

    document.add_heading("КЛЮЧЕВОЙ КАРЬЕРНЫЙ ЭФФЕКТ", level=2)
    _add_bullets(document, _career_impact_bullets(vacancy, language="ru"))

    document.add_heading("ОПЫТ", level=2)
    for heading, description, bullets in _resume_experience_blocks(vacancy, language="ru"):
        paragraph = document.add_paragraph()
        paragraph.add_run(heading).bold = True
        document.add_paragraph(description)
        key = document.add_paragraph()
        key.add_run("Ключевые достижения").bold = True
        _add_bullets(document, bullets)

    document.add_heading("ОБРАЗОВАНИЕ", level=2)
    document.add_paragraph("Информация об образовании предоставляется по запросу.")
    document.add_heading("ИНСТРУМЕНТЫ", level=2)
    document.add_paragraph("Figma; Miro; Jira; Confluence; Adobe Creative Suite; прототипирование; дизайн-системы; UX-исследования")
    document.add_heading("СЕРТИФИКАЦИИ", level=2)
    document.add_paragraph("Релевантные сертификаты в продуктовой, UX и leadership-практике предоставляются по запросу.")
    document.add_heading("ЯЗЫКИ", level=2)
    document.add_paragraph(profile.languages or "Русский: родной; английский: B2")
    return _save(document)


def _headline_tags(vacancy: Vacancy, language: str = "en") -> str:
    tags = _core_expertise(vacancy, language=language)[:5]
    return " • ".join(tags)


def _domain_focus(vacancy: Vacancy, language: str = "en") -> str:
    text = _vacancy_text(vacancy)
    if language == "ru":
        domains = ["продуктовая стратегия", "UX-стратегия", "дизайн-системы"]
        if _has_any(text, "ai", "llm", "prompt", "voice", "голос", "ии"):
            domains.append("AI и Voice UX")
        if _has_any(text, "automotive", "hmi", "vehicle", "авто", "электромоб"):
            domains.append("Automotive UX и HMI")
        if _has_any(text, "enterprise", "saas", "dashboard", "workflow", "b2b"):
            domains.append("Enterprise UX и сложные системы")
        return ", ".join(dict.fromkeys(domains))
    domains = ["Product Strategy", "UX Strategy", "Design Systems"]
    if _has_any(text, "ai", "llm", "prompt", "voice", "conversation", "intent"):
        domains.append("AI and Voice UX")
    if _has_any(text, "automotive", "hmi", "vehicle", "mobility", "infotainment"):
        domains.append("Automotive UX and HMI")
    if _has_any(text, "enterprise", "saas", "dashboard", "workflow", "b2b"):
        domains.append("Enterprise UX and complex systems")
    return ", ".join(dict.fromkeys(domains))


def _core_expertise(vacancy: Vacancy, language: str = "en") -> list[str]:
    text = _vacancy_text(vacancy)
    if language == "ru":
        base = _split_keywords(vacancy.top_match_keywords or vacancy.vacancy_keywords)[:10] + [
            "Продуктовая стратегия",
            "UX-стратегия",
            "Дизайн-лидерство",
            "Продуктовое исследование",
            "UX-исследования",
            "Кросс-функциональное взаимодействие",
            "Дизайн-системы",
            "Информационная архитектура",
        ]
        if _has_any(text, "ai", "llm", "prompt", "human-ai", "voice", "conversation", "intent", "голос", "ии"):
            base.extend(["AI-продуктовый дизайн", "Human-AI Interaction", "Voice UX", "Conversational Design", "Intent Design", "Мультимодальные интерфейсы"])
        if _has_any(text, "automotive", "hmi", "vehicle", "mobility", "infotainment", "avas", "авто", "электромоб"):
            base.extend(["Automotive UX", "HMI", "IVI", "Infotainment", "AVAS", "Опыт водителя", "Mobility"])
        if _has_any(text, "enterprise", "saas", "dashboard", "workflow", "operator", "b2b", "data"):
            base.extend(["Enterprise UX", "SaaS", "Дашборды", "Оптимизация workflow", "Сложные системы", "Визуализация данных"])
        if _has_any(text, "component", "tokens", "variables", "governance", "дизайн-систем"):
            base.extend(["Библиотека компонентов", "Design Tokens", "Figma Variables", "Design Governance", "DesignOps"])
        return list(dict.fromkeys(base))
    base = _split_keywords(vacancy.top_match_keywords or vacancy.vacancy_keywords)[:10] + [
        "Product Strategy",
        "UX Strategy",
        "Design Leadership",
        "Product Discovery",
        "User Research",
        "Cross-functional Collaboration",
        "Design Systems",
        "Information Architecture",
    ]
    if _has_any(text, "ai", "llm", "prompt", "human-ai", "voice", "conversation", "intent", "голос", "ии"):
        base.extend(["AI Product Design", "Human-AI Interaction", "Voice UX", "Conversational Design", "Intent Design", "Multimodal Interfaces"])
    if _has_any(text, "automotive", "hmi", "vehicle", "mobility", "infotainment", "avas", "авто", "электромоб"):
        base.extend(["Automotive UX", "HMI", "IVI", "Infotainment", "AVAS", "Driver Experience", "Mobility"])
    if _has_any(text, "enterprise", "saas", "dashboard", "workflow", "operator", "b2b", "data"):
        base.extend(["Enterprise UX", "SaaS", "Dashboard Design", "Workflow Optimization", "Complex Systems", "Data Visualization"])
    if _has_any(text, "component", "tokens", "variables", "governance", "дизайн-систем"):
        base.extend(["Component Library", "Design Tokens", "Figma Variables", "Design Governance", "DesignOps"])
    return list(dict.fromkeys(base))


def _career_impact_bullets(vacancy: Vacancy, language: str = "en") -> list[str]:
    text = _vacancy_text(vacancy)
    if language == "ru":
        bullets = [
            "Вел UX-стратегию для голосового ассистента электромобиля, определив conversational architecture, intents, prompts и принципы мультимодального взаимодействия.",
            "Сформировал продуктовую основу для автомобильных HMI-сценариев, связав customer journeys, системную логику интерфейса и требования инженерных команд.",
            "Укрепил дизайн-системный подход в B2B и B2C продуктах, выстроив повторно используемые паттерны, дизайн-ревью и принципы governance.",
            "Поставил enterprise UX-решения для smart city и транспортной инфраструктуры, упростив операторские сценарии для систем городского масштаба.",
        ]
        if _has_any(text, "voice", "conversation", "intent", "prompt", "голос"):
            bullets.insert(1, "Определил voice UX framework, который связал tone of voice, prompts, intents и визуальные состояния в единый пользовательский опыт.")
        if _has_any(text, "automotive", "hmi", "vehicle", "avas", "авто", "электромоб"):
            bullets.insert(2, "Вел продуктовую UX-логику звуковой идентичности и AVAS для production-ready электромобиля.")
        if _has_any(text, "component", "tokens", "design system", "дизайн-систем"):
            bullets.append("Масштабировал дизайн-системные практики через component library, Figma-подходы, governance и совместную работу с engineering.")
        return bullets[:7]
    bullets = [
        "Led UX strategy for the voice assistant powering an electric vehicle by defining conversational architecture, intents, prompts and multimodal interaction principles.",
        "Established a product foundation for automotive HMI experiences by connecting customer journeys, interface logic and engineering constraints.",
        "Scaled Design System practices across B2B and B2C products by introducing reusable patterns, design reviews and governance principles.",
        "Delivered enterprise UX solutions for smart city and transportation infrastructure by simplifying operator workflows for city-scale systems.",
    ]
    if _has_any(text, "voice", "conversation", "intent", "prompt"):
        bullets.insert(1, "Defined a Voice UX framework by aligning tone of voice, prompts, intents and visual states into a consistent product experience.")
    if _has_any(text, "automotive", "hmi", "vehicle", "avas"):
        bullets.insert(2, "Owned product UX logic for vehicle sound identity and AVAS by partnering with engineering and research teams from concept to production readiness.")
    if _has_any(text, "component", "tokens", "design system", "governance"):
        bullets.append("Accelerated product consistency by establishing component library principles, Figma workflows and design-system governance.")
    return bullets[:7]


def _resume_experience_blocks(vacancy: Vacancy, language: str = "en") -> list[tuple[str, str, list[str]]]:
    text = _vacancy_text(vacancy)
    if language == "ru":
        atom_bullets = [
            "Вел end-to-end продуктовый дизайн in-car голосового ассистента, определив conversational architecture, prompts, intents, tone of voice и мультимодальную UX-логику.",
            "Определил UX-принципы для HMI-сценариев электромобиля, связав customer journeys, продуктовые требования и инженерные ограничения в масштабируемую систему интерфейсных решений.",
            "Партнерски работал с Engineering, AI, Product и Research командами, чтобы объединить визуальное, голосовое и звуковое взаимодействие в цельный пользовательский опыт.",
        ]
        if _has_any(text, "design system", "component", "tokens", "дизайн-систем"):
            atom_bullets.append("Укрепил дизайн-системный подход для сложных automotive-сценариев, связав reusable patterns, interface states и продуктовую логику.")
        return [
            (
                "Lead Product Designer — ATOM, Electric Vehicle Startup",
                "Автомобильный продуктовый дизайн, voice UX, HMI, AI interaction и sound experience для электромобиля.",
                atom_bullets,
            ),
            (
                "Product Designer — Information Technology Factory",
                "Enterprise UX для smart city, транспорта, парковочных систем, dashboards и операторских интерфейсов.",
                [
                    "Поставил city-scale operator workflows, преобразовав сложные транспортные и парковочные процессы в понятные dashboards, mobile flows и интерфейсы мониторинга.",
                    "Оптимизировал мобильный опыт парковочного приложения, связав user research, information architecture, UX flows и production-ready интерфейсные решения.",
                ],
            ),
            (
                "Lead Product Designer — VEON / Beeline Russia",
                "Telecom-продукты B2B/B2C, дизайн-системы, UX-исследования и координация продуктовых дизайн-команд.",
                [
                    "Масштабировал продуктовый дизайн через design-system practices, reusable UI patterns, design reviews и взаимодействие с product и engineering командами.",
                    "Усилал delivery сложных telecom-продуктов, переводя неоднозначные требования в ясные user flows, prototypes и интерфейсные решения.",
                ],
            ),
            (
                "Product Designer — ABBYY Software House",
                "Международные software-продукты, сложные пользовательские сценарии и интерфейсная архитектура.",
                [
                    "Доставил UX/UI-решения для международных software-продуктов, структурировав product flows, interaction patterns и интерфейсную архитектуру для Lingvo и PDF Transformer.",
                ],
            ),
        ]

    atom_bullets = [
        "Led end-to-end product design of the in-vehicle voice assistant by defining conversational architecture, prompts, intents, tone of voice and multimodal UX logic.",
        "Defined UX principles for electric-vehicle HMI scenarios by connecting customer journeys, product requirements and engineering constraints into a scalable interaction system.",
        "Partnered with Engineering, AI, Product and Research teams to deliver multimodal user experiences combining visual, voice and audio interaction.",
    ]
    if _has_any(text, "design system", "component", "tokens", "governance"):
        atom_bullets.append("Established design-system principles for complex automotive scenarios by aligning reusable patterns, interface states and product logic.")
    return [
        (
            "Lead Product Designer — ATOM, Electric Vehicle Startup",
            "Automotive product design, Voice UX, HMI, AI interaction and sound experience for an electric vehicle.",
            atom_bullets,
        ),
        (
            "Product Designer — Information Technology Factory",
            "Enterprise UX for smart city, transportation, parking systems, dashboards and operator interfaces.",
            [
                "Delivered city-scale operator workflows by transforming complex transportation and parking processes into clear dashboards, mobile flows and monitoring interfaces.",
                "Optimized the mobile parking experience by connecting user research, information architecture, UX flows and production-ready interface solutions.",
            ],
        ),
        (
            "Lead Product Designer — VEON / Beeline Russia",
            "B2B/B2C telecom products, design systems, UX research and product design team coordination.",
            [
                "Scaled product design through design-system practices, reusable UI patterns, design reviews and close partnership with product and engineering teams.",
                "Transformed ambiguous telecom requirements into clear user flows, prototypes and interface solutions for complex B2B and B2C products.",
            ],
        ),
        (
            "Product Designer — ABBYY Software House",
            "International software products, complex user scenarios and interface architecture.",
            [
                "Delivered UX/UI solutions for international software products by structuring product flows, interaction patterns and interface architecture for Lingvo and PDF Transformer.",
            ],
        ),
    ]


def _vacancy_text(vacancy: Vacancy) -> str:
    return " ".join(
        [
            vacancy.title,
            vacancy.company,
            vacancy.description_raw,
            vacancy.requirements,
            vacancy.responsibilities,
            vacancy.vacancy_keywords,
            vacancy.top_match_keywords,
        ]
    ).lower()


def _has_any(text: str, *terms: str) -> bool:
    return any(term in text for term in terms)


def _base_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(9.5)
    return document


def _candidate_display_name(profile: CandidateProfile) -> str:
    name = profile.name.strip()
    if not name or name.lower() == "candidate":
        return "ALEKSANDR GRENKOV"
    return name.upper()


def _add_title(
    document: Document,
    name: str,
    headline: str,
    language: str = "en",
    profile: Optional[CandidateProfile] = None,
) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(name)
    run.bold = True
    run.font.size = Pt(16)
    subtitle = document.add_paragraph()
    subtitle.add_run(headline).bold = True
    document.add_paragraph(_contact_line(profile))


def _contact_line(profile: Optional[CandidateProfile]) -> str:
    contacts = _extract_contacts(profile.raw_cv_text if profile else "")
    parts = [
        contacts.get("email") or "Email",
        contacts.get("linkedin") or "LinkedIn",
        contacts.get("portfolio") or "Portfolio",
        contacts.get("telegram") or "Telegram",
        "Open to international and remote opportunities.",
    ]
    return " • ".join(parts)


def _extract_contacts(text: str) -> dict[str, str]:
    urls = [_clean_contact(match) for match in re.findall(r"https?://[^\s,)>\]]+", text)]
    email_match = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)
    telegram_match = re.search(r"(?:^|[\s:])(@[A-Za-z0-9_]{4,})\b", text, flags=re.IGNORECASE)
    linkedin = next((url for url in urls if "linkedin.com/" in url.lower()), "")
    telegram_url = next((url for url in urls if "t.me/" in url.lower()), "")
    portfolio = next(
        (
            url
            for url in urls
            if "linkedin.com/" not in url.lower() and "t.me/" not in url.lower() and "hh.ru/" not in url.lower()
        ),
        "",
    )
    return {
        "email": email_match.group(0) if email_match else "",
        "linkedin": linkedin,
        "portfolio": portfolio,
        "telegram": telegram_match.group(1) if telegram_match else telegram_url,
    }


def _clean_contact(value: str) -> str:
    return value.strip().rstrip(".,;:")


def _add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        if item.strip():
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(item.strip())


def _add_tag_paragraph(document: Document, items: list[str]) -> None:
    tags = [item.strip() for item in items if item.strip()]
    if tags:
        document.add_paragraph(" • ".join(tags))


def _add_experience_template(document: Document, heading: str, bullets: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(heading).bold = True
    _add_bullets(document, bullets)


def _split_keywords(value: str) -> list[str]:
    return [item.strip() for item in re.split(r"[;\n,|•]+", value or "") if item.strip()]


def _is_russian_vacancy(vacancy: Vacancy) -> bool:
    language = vacancy.language.lower()
    text = " ".join([vacancy.title, vacancy.description_raw, vacancy.requirements, vacancy.responsibilities])
    cyrillic = sum(1 for char in text if "а" <= char.lower() <= "я" or char.lower() == "ё")
    latin = sum(1 for char in text if "a" <= char.lower() <= "z")
    if cyrillic or latin:
        return cyrillic >= 12 and cyrillic > latin * 0.25
    if "russian" in language and "english" not in language:
        return True
    return False


def _save(document: Document) -> BytesIO:
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
