import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.models import ApplicationMaterial, Vacancy
from app.services.assets import pdf_path_for_cv, png_preview_path_for_cv
from app.status import ApplicationStatus


READY_STATUS_ALIASES = {"Ready to submit manually", "Ready to send"}


def _value(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _int_value(value: Any, default: int = 0) -> int:
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return default


def import_csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8") as file:
        return list(csv.DictReader(file))


def import_xlsx_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[sheet_name]
    headers = [_value(cell.value) for cell in sheet[1]]
    rows: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
        if any(_value(value) for value in record.values()):
            rows.append(record)
    return rows


def normalize_queue_row(row: dict[str, Any], package_root: Path) -> Vacancy:
    cv_file_path = _value(row.get("CV File Path"))
    submit_status = _value(row.get("Submit Status"))
    status = ApplicationStatus.READY_TO_SEND if submit_status in READY_STATUS_ALIASES else ApplicationStatus.DRAFT
    return Vacancy(
        external_id=_value(row.get("ID")),
        rank=_int_value(row.get("Rank"), default=0) or None,
        source=_value(row.get("Source")),
        company=_value(row.get("Company")),
        title=_value(row.get("Vacancy")),
        location=_value(row.get("Location")),
        posted=_value(row.get("Posted")),
        date_status=_value(row.get("Date Status")),
        language=_value(row.get("Language")),
        fit_score=_int_value(row.get("Fit")),
        priority=_value(row.get("Priority")) or "Medium",
        submit_status=status,
        next_action=_value(row.get("Next Action")),
        cv_file_path=cv_file_path,
        pdf_file_path=pdf_path_for_cv(cv_file_path),
        png_preview_path=png_preview_path_for_cv(cv_file_path),
        source_url=_value(row.get("Source URL")) or _value(row.get("Vacancy Link")),
        tailored_headline=_value(row.get("Tailored Headline")),
        top_match_keywords=_value(row.get("Top Match Keywords")),
        gaps_risks=_value(row.get("Gaps / Risks")),
        adaptation_strategy=_value(row.get("Adaptation Strategy")),
    )


def normalize_material_row(row: dict[str, Any], vacancy_id: int) -> ApplicationMaterial:
    fit_summary = (
        _value(row.get("Fit Summary")) or _value(row.get("Summary")) or "Imported from tailored package."
    )
    return ApplicationMaterial(
        vacancy_id=vacancy_id,
        short_note=_value(row.get("Short Platform Note")),
        recruiter_dm=_value(row.get("Recruiter / Telegram DM")),
        email_cover_letter=_value(row.get("Email Cover Letter")),
        fit_summary=fit_summary,
    )
